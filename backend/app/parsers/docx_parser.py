import io
from typing import Dict, List, Tuple
from docx import Document

class DOCXParsingError(Exception):
    pass

def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, List[Dict[str, any]]]:
    """
    Extracts text from a DOCX file using python-docx.
    Returns:
        - formatted_text: Structured document text with paragraph/heading tags.
        - sections: List of detected sections and paragraph blocks.
    Raises:
        - DOCXParsingError: If file is corrupted or not a valid Word document.
    """
    try:
        doc_stream = io.BytesIO(file_bytes)
        doc = Document(doc_stream)
    except Exception as e:
        raise DOCXParsingError(f"Could not parse DOCX file. Ensure it is a valid .docx document: {str(e)}")

    sections = []
    lines = []
    current_section = "Preamble"

    # Extract paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Check if heading
        if para.style and para.style.name and "Heading" in para.style.name:
            current_section = text
            lines.append(f"\n### {text} ###")
            sections.append({"section": text, "type": "heading"})
        else:
            lines.append(text)
            sections.append({"section": current_section, "text": text, "type": "paragraph"})

    # Extract tables if any
    for table_idx, table in enumerate(doc.tables):
        lines.append(f"\n[Table {table_idx + 1}]")
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                lines.append(row_text)

    formatted_text = "\n\n".join(lines).strip()

    if not formatted_text or len(formatted_text) < 10:
        raise DOCXParsingError("The DOCX document contains no readable text.")

    return formatted_text, sections
