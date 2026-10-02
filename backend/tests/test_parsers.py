import io
import pytest
import pymupdf as fitz
from docx import Document
from app.parsers.pdf_parser import extract_text_from_pdf, PDFParsingError
from app.parsers.docx_parser import extract_text_from_docx, DOCXParsingError
from app.parsers.text_parser import extract_text_from_raw, TextParsingError

def test_pdf_parser_success():
    # Create an in-memory PDF with 2 pages
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((50, 50), "MASTER SERVICES AGREEMENT\nEffective Date: January 15, 2026\nParties: Acme Corp and Beta LLC.")
    p2 = doc.new_page()
    p2.insert_text((50, 50), "SECTION 4. TERMINATION\nEither party may terminate upon 30 days written notice.")
    
    pdf_bytes = doc.tobytes()
    doc.close()

    text, page_data = extract_text_from_pdf(pdf_bytes)
    assert "[Page 1]" in text
    assert "[Page 2]" in text
    assert "MASTER SERVICES AGREEMENT" in text
    assert "Either party may terminate upon 30 days" in text
    assert len(page_data) == 2
    assert page_data[0]["page"] == 1
    assert page_data[1]["page"] == 2

def test_pdf_parser_empty_or_scanned():
    # Empty PDF with 0 pages or blank page
    doc = fitz.open()
    doc.new_page() # blank page without text
    pdf_bytes = doc.tobytes()
    doc.close()

    with pytest.raises(PDFParsingError) as exc_info:
        extract_text_from_pdf(pdf_bytes)
    assert "No extractable text found" in str(exc_info.value)

def test_docx_parser_success():
    doc = Document()
    doc.add_heading("SOFTWARE LICENSE AGREEMENT", level=1)
    doc.add_paragraph("This Agreement is entered into by and between Alpha Inc and Omega Corp.")
    doc.add_heading("Section 2: Term and Renewal", level=2)
    doc.add_paragraph("This agreement automatically renews annually unless 60 days notice is provided.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Date"
    table.cell(1, 0).text = "Delivery"
    table.cell(1, 1).text = "2026-10-31"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    text, sections = extract_text_from_docx(docx_bytes)
    assert "SOFTWARE LICENSE AGREEMENT" in text
    assert "Section 2: Term and Renewal" in text
    assert "Milestone | Date" in text
    assert "Delivery | 2026-10-31" in text
    assert len(sections) > 0

def test_docx_parser_corrupted():
    with pytest.raises(DOCXParsingError):
        extract_text_from_docx(b"not a docx file")

def test_text_parser_success():
    raw = """
    NON-DISCLOSURE AGREEMENT
    Effective Date: 2026-03-01
    Section 1. Confidentiality
    The Receiving Party shall maintain confidentiality for 3 years.
    Section 2. Return of Materials
    Upon termination, materials must be destroyed within 10 days.
    """
    cleaned, sections = extract_text_from_raw(raw)
    assert "NON-DISCLOSURE AGREEMENT" in cleaned
    assert len(sections) >= 2

def test_text_parser_empty():
    with pytest.raises(TextParsingError):
        extract_text_from_raw("   ")
